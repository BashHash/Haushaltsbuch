from __future__ import annotations

import csv
import json
import shutil
import sqlite3
import sys
import uuid
from datetime import datetime as dt_datetime
from collections import defaultdict
from dataclasses import dataclass, asdict
from datetime import date, datetime
from pathlib import Path

from PySide6.QtCore import (QAbstractTableModel, QModelIndex, Qt, QDate,
                            QSortFilterProxyModel, Signal)
from PySide6.QtGui import (QAction, QColor, QPainter, QPen, QIcon,
                           QGuiApplication)
from PySide6.QtWidgets import (
    QApplication, QCheckBox, QComboBox, QDateEdit, QDialog, QDialogButtonBox,
    QDoubleSpinBox, QFileDialog, QFormLayout, QHBoxLayout, QHeaderView, QLabel,
    QLineEdit, QListWidget, QMainWindow, QMessageBox, QPushButton, QSpinBox,
    QSplitter, QStyleFactory, QTabWidget, QTableView, QTextEdit, QVBoxLayout,
    QWidget, QAbstractItemView, QGroupBox
)
try:
    from PySide6.QtCharts import (QChart, QChartView, QLineSeries, QBarSeries,
                                  QBarSet, QBarCategoryAxis, QValueAxis,
                                  QStackedBarSeries)
    CHARTS_AVAILABLE = True
except ImportError:
    CHARTS_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill  # ,alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

VERSION = "1.1.0"
DEVELOPER = "Andreas Windmüller\nawindi@gmx.de"
APP_TITLE = "Haushaltsbuch - V" + VERSION
DATE_FMT = "%d.%m.%Y"
DEFAULT_FILE = Path.home() / "haushaltsbuch.json"
RECENT_FILES_FILE = Path.home() / ".haushaltsbuch_recent.json"
MAX_RECENT_FILES = 3


def money(value: float) -> str:
    s = f"{value:,.2f} €"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def parse_bool(value) -> bool:
    return str(value).strip().lower() in {"1", "true", "ja", "yes", "y", "x"}


@dataclass
class Booking:
    id: str
    datum: str
    kategorie: str
    beschreibung: str
    paypal: bool
    regelmaessig: bool
    konto: str
    betrag: float


class Store:
    def __init__(self):
        self.path: Path | None = None
        self.categories = ["Abhebung", "Auto",
                           "Bank", "Darlehen",
                           "Dienstleistung", "Einkaufen",
                           "Gehalt", "Internetshop",
                           "Medien", "Mitgliedsbeitrag",
                           "Spende", "Telefon",
                           "Urlaub", "Versicherung",
                           "Wohnen", "Sonstiges"]
        self.descriptions = ["Arzt xyz", "Bank xyz", "Bus/Bahn",
                             "Darlehen 12345", "Gebühren",
                             "Kartenpreis", "Kontoabschluss",
                             "Kreditkarte", "Mautgebühr",
                             "Rundfunkgebühr", "Taxi",
                             "Tankstelle xyz", "Umbuchung",
                             "Versicherung xyz"]
        self.accounts = [{"name": "Girokonto", "opening_balance": 0.0}]
        self.bookings: list[Booking] = []

    def to_dict(self):
        return {
            "version": 1,
            "settings": {
                "categories": self.categories,
                "descriptions": self.descriptions,
                "accounts": self.accounts,
            },
            "bookings": [asdict(x) for x in self.bookings],
        }

    def _create_backup(self):
        """Create a timestamped backup before overwriting."""
        if self.path and self.path.exists():
            backup_dir = self.path.parent / "backups"
            backup_dir.mkdir(exist_ok=True)
            timestamp = dt_datetime.now().strftime("%Y%m%d_%H%M%S")
            backup_path = backup_dir / \
                f"{self.path.stem}_{timestamp}{self.path.suffix}"
            shutil.copy2(self.path, backup_path)
            # Keep only the last 10 backups
            backups = sorted(backup_dir.glob(
                f"{self.path.stem}_*{self.path.suffix}"))
            for old in backups[:-10]:
                old.unlink(missing_ok=True)

    def save(self, path: Path | None = None):
        if path:
            self.path = path
        if not self.path:
            self.path = DEFAULT_FILE
        self._create_backup()
        self.path.write_text(json.dumps(self.to_dict(),
                                        ensure_ascii=False,
                                        indent=2),
                             encoding="utf-8")

    def load(self, path: Path):
        raw = json.loads(path.read_text(encoding="utf-8"))
        settings = raw.get("settings", {})
        self.categories = list(settings.get("categories", []))
        self.descriptions = list(settings.get("descriptions", []))
        self.accounts = list(settings.get("accounts", []))
        self.bookings = [Booking(**row) for row in raw.get("bookings", [])]
        self.path = path

    # --- SQLite support ---

    def save_sqlite(self, path: Path):
        """Save all data to a SQLite database."""
        conn = sqlite3.connect(str(path))
        cur = conn.cursor()
        cur.executescript("""
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT
            );
            CREATE TABLE IF NOT EXISTS accounts (
                name TEXT PRIMARY KEY,
                opening_balance REAL
            );
            CREATE TABLE IF NOT EXISTS bookings (
                id TEXT PRIMARY KEY,
                datum TEXT,
                kategorie TEXT,
                beschreibung TEXT,
                paypal INTEGER,
                regelmaessig INTEGER,
                konto TEXT,
                betrag REAL
            );
        """)
        # Clear existing data
        cur.execute("DELETE FROM settings")
        cur.execute("DELETE FROM accounts")
        cur.execute("DELETE FROM bookings")
        # Save settings
        cur.execute("INSERT INTO settings VALUES (?, ?)",
                    ("categories", json.dumps(self.categories)))
        cur.execute("INSERT INTO settings VALUES (?, ?)",
                    ("descriptions", json.dumps(self.descriptions)))
        # Save accounts
        for a in self.accounts:
            cur.execute("INSERT INTO accounts VALUES (?, ?)",
                        (a["name"], a.get("opening_balance", 0.0)))
        # Save bookings
        for b in self.bookings:
            cur.execute("""INSERT INTO bookings VALUES (?,?,?,?,?,?,?,?)""",
                        (b.id, b.datum, b.kategorie, b.beschreibung,
                         int(b.paypal), int(b.regelmaessig),
                         b.konto, b.betrag))
        conn.commit()
        conn.close()
        self.path = path

    def load_sqlite(self, path: Path):
        """Load all data from a SQLite database."""
        conn = sqlite3.connect(str(path))
        cur = conn.cursor()
        # Load settings
        cur.execute("SELECT key, value FROM settings")
        settings = dict(cur.fetchall())
        self.categories = json.loads(settings.get("categories", "[]"))
        self.descriptions = json.loads(settings.get("descriptions", "[]"))
        # Load accounts
        cur.execute("SELECT name, opening_balance FROM accounts")
        self.accounts = [{"name": r[0],
                          "opening_balance": r[1]} for r in cur.fetchall()]
        # Load bookings
        cur.execute("""SELECT id, datum, kategorie, beschreibung, paypal,
                              regelmaessig, konto, betrag FROM bookings""")
        self.bookings = [Booking(id=r[0], datum=r[1], kategorie=r[2],
                                 beschreibung=r[3], paypal=bool(r[4]),
                                 regelmaessig=bool(r[5]), konto=r[6],
                                 betrag=r[7]) for r in cur.fetchall()]
        conn.close()
        self.path = path

    # --- Excel export ---

    def export_excel(self, path: Path):
        """Export bookings and settings to an Excel file."""
        if not EXCEL_AVAILABLE:
            raise RuntimeError("openpyxl ist nicht installiert.")
        wb = openpyxl.Workbook()

        # Bookings sheet
        ws = wb.active
        ws.title = "Buchungen"
        headers = ["Datum", "Kategorie", "Beschreibung", "PayPal",
                   "Regelmäßig", "Konto", "Betrag"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E2F3")
        for b in self.bookings:
            ws.append([format_date_german(b.datum), b.kategorie,
                       b.beschreibung,
                       "Ja" if b.paypal else "Nein",
                       "Ja" if b.regelmaessig else "Nein",
                       b.konto, b.betrag])
        # Format amount column
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            for cell in row:
                cell.number_format = '#,##0.00 €'
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['G'].width = 14

        # Accounts sheet
        ws2 = wb.create_sheet("Konten")
        ws2.append(["Konto", "Anfangsstand"])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E2F3")
        for a in self.accounts:
            ws2.append([a["name"], a.get("opening_balance", 0.0)])
        for row in ws2.iter_rows(min_row=2, min_col=2, max_col=2):
            for cell in row:
                cell.number_format = '#,##0.00 €'

        # Settings sheet
        ws3 = wb.create_sheet("Einstellungen")
        ws3.append(["Kategorien"])
        for c in self.categories:
            ws3.append([c])
        ws3.append([])
        ws3.append(["Beschreibungen"])
        for d in self.descriptions:
            ws3.append([d])

        wb.save(str(path))


class BookingModel(QAbstractTableModel):
    headers = ["Datum", "Kategorie", "Beschreibung", "PayPal",
               "Regelmäßig", "Konto", "Betrag"]

    def __init__(self, store: Store):
        super().__init__()
        self.store = store

    def rowCount(self, parent=QModelIndex()):
        return len(self.store.bookings)

    def columnCount(self, parent=QModelIndex()):
        return len(self.headers)

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role == Qt.DisplayRole and orientation == Qt.Horizontal:
            return self.headers[section]
        return None

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        b = self.store.bookings[index.row()]
        values = [b.datum, b.kategorie, b.beschreibung, b.paypal,
                  b.regelmaessig, b.konto, b.betrag]
        v = values[index.column()]
        if role == Qt.DisplayRole:
            if index.column() in (3, 4):
                return "Ja" if v else "Nein"
            if index.column() == 6:
                return money(float(v))
            if index.column() == 0:
                return format_date_german(v)
            return str(v)
        if role == Qt.UserRole:
            return v
        if role == Qt.TextAlignmentRole and index.column() == 6:
            return int(Qt.AlignRight | Qt.AlignVCenter)
        if role == Qt.ForegroundRole and index.column() == 6:
            return QColor("#0b7a38") if float(v) >= 0 else QColor("#b42318")
        return None

    def refresh(self):
        self.beginResetModel()
        self.endResetModel()


def format_date_german(iso_date: str) -> str:
    """Convert ISO date string (yyyy-MM-dd) to German format (dd.MM.yyyy)."""
    try:
        return datetime.strptime(str(iso_date), "%Y-%m-%d").strftime(DATE_FMT)
    except (ValueError, TypeError):
        return str(iso_date)


class BookingSortProxy(QSortFilterProxyModel):
    def data(self, index, role=Qt.DisplayRole):
        if role == Qt.DisplayRole and index.column() == 0:
            source_index = self.mapToSource(index)
            if source_index.isValid():
                source_data = self.sourceModel().data(source_index,
                                                      Qt.UserRole)
                return format_date_german(source_data)
        return super().data(index, role)

    def lessThan(self, left, right):
        lv = left.data(Qt.UserRole)
        rv = right.data(Qt.UserRole)
        if left.column() == 0:
            try:
                # UserRole stores ISO format (yyyy-MM-dd)
                return datetime.strptime(lv, "%Y-%m-%d") \
                    < datetime.strptime(rv, "%Y-%m-%d")
            except Exception:
                pass
        if left.column() == 6:
            return float(lv) < float(rv)
        return str(lv).lower() < str(rv).lower()


class MultiFilterProxy(BookingSortProxy):
    def __init__(self):
        super().__init__()
        self.year = 0
        self.month = 0
        self.category = ""
        self.account = ""
        self.description = ""

    def set_filters(self, year=0, month=0, category="",
                    account="", description=""):
        self.year, self.month = year, month
        self.category, self.account, self.description = \
            category, account, description.strip().lower()
        self.beginFilterChange()
        self.endFilterChange()

    def filterAcceptsRow(self, row, parent):
        def get(col):
            return m.index(row, col, parent).data(Qt.UserRole)
        m = self.sourceModel()
        # get = lambda col: m.index(row, col, parent).data(Qt.UserRole)
        try:
            # UserRole stores ISO format (yyyy-MM-dd)
            d = datetime.strptime(str(get(0)), "%Y-%m-%d")
        except ValueError:
            return False
        return (
            (not self.year or d.year == self.year)
            and (not self.month or d.month == self.month)
            and (not self.category or get(1) == self.category)
            and (not self.account or get(5) == self.account)
            and (not self.description or self.description in
                 str(get(2)).lower())
        )


class BookingDialog(QDialog):
    def __init__(self, store: Store, booking: Booking | None = None,
                 parent=None):
        super().__init__(parent)
        self.setWindowTitle("Buchung bearbeiten")
        self.store, self.booking = store, booking
        form = QFormLayout(self)
        self.date = QDateEdit(calendarPopup=True)
        self.date.setDisplayFormat("dd.MM.yyyy")
        self.category = QComboBox()
        self.category.addItems(store.categories)
        self.description = QComboBox()
        self.description.setEditable(True)
        self.description.addItems(store.descriptions)
        self.paypal = QCheckBox("Ja")
        self.regular = QCheckBox("Ja")
        self.account = QComboBox()
        self.account.addItems([x["name"] for x in store.accounts])
        self.amount = QDoubleSpinBox()
        self.amount.setRange(-999999999, 999999999)
        self.amount.setDecimals(2)
        self.amount.setSuffix(" €")
        for label, widget in [("Datum", self.date),
                              ("Kategorie", self.category),
                              ("Beschreibung", self.description),
                              ("PayPal-Buchung", self.paypal),
                              ("Regelmäßige Buchung", self.regular),
                              ("Konto", self.account),
                              ("Betrag", self.amount)]:
            form.addRow(label, widget)
        buttons = QDialogButtonBox(QDialogButtonBox.Save |
                                   QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        form.addRow(buttons)
        if booking:
            self.date.setDate(QDate.fromString(booking.datum, "yyyy-MM-dd"))
            self.category.setCurrentText(booking.kategorie)
            self.description.setCurrentText(booking.beschreibung)
            self.paypal.setChecked(booking.paypal)
            self.regular.setChecked(booking.regelmaessig)
            self.account.setCurrentText(booking.konto)
            self.amount.setValue(booking.betrag)
        else:
            self.date.setDate(QDate.currentDate())

    def result_booking(self):
        return Booking(
            id=self.booking.id if self.booking else str(uuid.uuid4()),
            datum=self.date.date().toString("yyyy-MM-dd"),
            kategorie=self.category.currentText(),
            beschreibung=self.description.currentText().strip(),
            paypal=self.paypal.isChecked(),
            regelmaessig=self.regular.isChecked(),
            konto=self.account.currentText(),
            betrag=self.amount.value())


class EntryTab(QWidget):
    changed = Signal()

    def __init__(self, store, model):
        super().__init__()
        self.store, self.model = store, model
        root = QVBoxLayout(self)
        box = QGroupBox("Neue Buchung")
        form = QHBoxLayout(box)
        self.date = QDateEdit(QDate.currentDate())
        self.date.setCalendarPopup(True)
        self.date.setDisplayFormat("dd.MM.yyyy")
        self.category = QComboBox()
        self.description = QComboBox()
        self.description.setEditable(True)
        self.paypal = QCheckBox("PayPal")
        self.regular = QCheckBox("Regelmäßig")
        self.account = QComboBox()
        self.amount = QDoubleSpinBox()
        self.amount.setRange(-999999999, 999999999)
        self.amount.setDecimals(2)
        self.amount.setSuffix(" €")
        self.add_btn = QPushButton("Hinzufügen")

        for label, w in [("Datum", self.date),
                         ("Kategorie", self.category),
                         ("Beschreibung", self.description),
                         ("", self.paypal),
                         ("", self.regular),
                         ("Konto", self.account),
                         ("Betrag", self.amount),
                         ("", self.add_btn)]:
            cell = QVBoxLayout()
            cell.addWidget(QLabel(label))
            cell.addWidget(w)
            form.addLayout(cell)
        root.addWidget(box)

        # --- Filter (Jahr, Monat, Konto) ---
        filter_box = QGroupBox("Filter")
        filter_layout = QHBoxLayout(filter_box)
        self.year_filter = QComboBox()
        self.month_filter = QComboBox()
        self.account_filter = QComboBox()
        self.month_filter.addItem("Alle Monate", 0)
        for i, n in enumerate(["Januar", "Februar", "März", "April", "Mai",
                               "Juni", "Juli", "August", "September",
                               "Oktober", "November", "Dezember"], 1):
            self.month_filter.addItem(n, i)
        for label, w in [("Jahr", self.year_filter),
                         ("Monat", self.month_filter),
                         ("Konto", self.account_filter)]:
            cell = QVBoxLayout()
            cell.addWidget(QLabel(label))
            cell.addWidget(w)
            filter_layout.addLayout(cell)
        filter_layout.addStretch()
        root.addWidget(filter_box)

        self.proxy = MultiFilterProxy()
        self.proxy.setSourceModel(model)
        self.table = QTableView()
        self.table.setModel(self.proxy)
        self.table.setSortingEnabled(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        root.addWidget(self.table)

        row = QHBoxLayout()
        edit = QPushButton("Ausgewählte bearbeiten")
        delete = QPushButton("Ausgewählte löschen")
        row.addWidget(edit)
        row.addWidget(delete)
        row.addStretch()
        root.addLayout(row)

        # --- Kontostände ---
        balance_box = QGroupBox("Kontostände")
        balance_layout = QVBoxLayout(balance_box)
        self.balance_table = QTableView()
        self.balance_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.Stretch)
        balance_layout.addWidget(self.balance_table)
        root.addWidget(balance_box)

        self.add_btn.clicked.connect(self.add)
        self.amount.lineEdit().returnPressed.connect(self.add)
        edit.clicked.connect(self.edit)
        delete.clicked.connect(self.delete)
        self.table.doubleClicked.connect(lambda _=None: self.edit())

        for w in (self.year_filter, self.month_filter, self.account_filter):
            w.currentIndexChanged.connect(self.apply_filters)
        self.refresh_settings()

    def refresh_settings(self):
        cat, acc = self.category.currentText(), self.account.currentText()
        self.category.clear()
        self.category.addItems(self.store.categories)
        self.description.clear()
        self.description.addItems(self.store.descriptions)
        self.account.clear()
        self.account.addItems([x["name"] for x in self.store.accounts])
        if cat:
            self.category.setCurrentText(cat)
        if acc:
            self.account.setCurrentText(acc)

        # Refresh filter combos
        years = sorted({int(x.datum[:4]) for x in self.store.bookings
                        if len(x.datum) >= 4}, reverse=True)
        self.year_filter.blockSignals(True)
        self.year_filter.clear()
        self.year_filter.addItem("Alle Jahre", 0)
        for y in years:
            self.year_filter.addItem(str(y), y)
        self.year_filter.blockSignals(False)

        acc_filter = self.account_filter.currentData()
        self.account_filter.blockSignals(True)
        self.account_filter.clear()
        self.account_filter.addItem("Alle Konten", "")
        for a in self.store.accounts:
            self.account_filter.addItem(a["name"], a["name"])
        self.account_filter.blockSignals(False)
        if acc_filter is not None:
            idx = self.account_filter.findData(acc_filter)
            if idx >= 0:
                self.account_filter.setCurrentIndex(idx)

        self.apply_filters()

    def apply_filters(self):
        year = self.year_filter.currentData() or 0
        month = self.month_filter.currentData() or 0
        account = self.account_filter.currentData() or ""
        self.proxy.set_filters(year, month, "", account, "")
        self._refresh_balances()

    def _refresh_balances(self):
        year = self.year_filter.currentData() or 0
        month = self.month_filter.currentData() or 0
        account_filter = self.account_filter.currentData() or ""
        rows = []
        total = 0.0
        for a in self.store.accounts:
            name = a["name"]
            if account_filter and name != account_filter:
                continue
            start = float(a.get("opening_balance", 0))
            s = 0.0
            for b in self.store.bookings:
                if b.konto != name:
                    continue
                if year:
                    try:
                        by = int(b.datum[:4])
                    except (ValueError, TypeError):
                        continue
                    if by > year:
                        continue
                    if by == year and month:
                        try:
                            bm = int(b.datum[5:7])
                        except (ValueError, TypeError):
                            continue
                        if bm > month:
                            continue
                s += b.betrag
            end = start + s
            total += end
            rows.append([name, money(start), money(end)])
        rows.append(["Summe", "", money(total)])
        self.balance_table.setModel(SimpleTableModel(
            ["Konto", "Anfangsstand", "Kontostand"], rows))

    def add(self):
        if not self.category.currentText() or not self.account.currentText():
            return
        b = Booking(str(uuid.uuid4()),
                    self.date.date().toString("yyyy-MM-dd"),
                    self.category.currentText(),
                    self.description.currentText().strip(),
                    self.paypal.isChecked(),
                    self.regular.isChecked(),
                    self.account.currentText(),
                    self.amount.value())
        row = len(self.store.bookings)
        self.table.setUpdatesEnabled(False)
        self.model.beginInsertRows(QModelIndex(), row, row)
        self.store.bookings.append(b)
        self.model.endInsertRows()
        self.table.setUpdatesEnabled(True)
        self._refresh_balances()
        self.amount.setValue(0)
        self.description.setCurrentText("")
        self.paypal.setChecked(False)
        self.regular.setChecked(False)
        self.date.setFocus()
        self.changed.emit()

    def selected_source_row(self):
        index = self.table.currentIndex()
        if not index.isValid():
            return -1
        source_index = self.proxy.mapToSource(index)
        return source_index.row()

    def edit(self):
        row = self.selected_source_row()
        if row < 0:
            return
        dlg = BookingDialog(self.store,
                            self.store.bookings[row],
                            self)
        if dlg.exec():
            self.store.bookings[row] = dlg.result_booking()
            self.model.refresh()
            self._refresh_balances()
            self.changed.emit()

    def delete(self):
        indexes = self.table.selectionModel().selectedRows()

        if not indexes:
            return

        rows = sorted(
            [self.proxy.mapToSource(idx).row() for idx in indexes],
            reverse=True
            )

        if QMessageBox.question(self,
                                "Löschen",
                                f"{len(rows)} Buchung(en) löschen?"
                                ) != QMessageBox.Yes:
            return

        for row in rows:
            del self.store.bookings[row]

        self.model.refresh()
        self._refresh_balances()
        self.changed.emit()


class PivotTab(QWidget):
    def __init__(self, store, model):
        super().__init__()
        self.store = store
        root = QVBoxLayout(self)
        filters = QHBoxLayout()
        self.year = QComboBox()
        self.month = QComboBox()
        self.category = QComboBox()
        self.account = QComboBox()
        self.description = QLineEdit()
        self.description.setPlaceholderText("enthält …")
        self.month.addItem("Alle Monate", 0)

        for i, n in enumerate(["Januar",
                               "Februar",
                               "März",
                               "April",
                               "Mai",
                               "Juni",
                               "Juli",
                               "August",
                               "September",
                               "Oktober",
                               "November",
                               "Dezember"],
                              1):
            self.month.addItem(n, i)
        for label, w in [("Jahr", self.year),
                         ("Monat", self.month),
                         ("Kategorie", self.category),
                         ("Konto", self.account),
                         ("Beschreibung", self.description)]:
            cell = QVBoxLayout()
            cell.addWidget(QLabel(label))
            cell.addWidget(w)
            filters.addLayout(cell)
        root.addLayout(filters)

        splitter = QSplitter(Qt.Vertical)
        self.proxy = MultiFilterProxy()
        self.proxy.setSourceModel(model)
        self.table = QTableView()
        self.table.setModel(self.proxy)
        self.table.setSortingEnabled(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.summary = QTableView()
        splitter.addWidget(self.table)
        splitter.addWidget(self.summary)
        root.addWidget(splitter)

        for w in (self.year,
                  self.month,
                  self.category,
                  self.account):
            w.currentIndexChanged.connect(self.apply)
        self.description.textChanged.connect(self.apply)
        self.refresh_settings()
        self.apply()

    def refresh_settings(self):
        years = sorted({int(x.datum[:4]) for x in self.store.bookings
                        if len(x.datum) >= 4}, reverse=True)
        self.year.blockSignals(True)
        self.year.clear()
        self.year.addItem("Alle Jahre", 0)
        for y in years:
            self.year.addItem(str(y), y)
        self.year.blockSignals(False)

        for combo, vals, label in [(self.category, self.store.categories,
                                    "Alle Kategorien"),
                                   (self.account, [x["name"] for x in
                                                   self.store.accounts],
                                    "Alle Konten")]:
            combo.blockSignals(True)
            combo.clear()
            combo.addItem(label, "")
            [combo.addItem(v, v) for v in vals]
            combo.blockSignals(False)

    def apply(self):
        self.proxy.set_filters(self.year.currentData() or
                               0, self.month.currentData() or
                               0, self.category.currentData() or
                               "", self.account.currentData() or
                               "", self.description.text())
        sums = defaultdict(float)
        for r in range(self.proxy.rowCount()):
            sums[self.proxy.index(r, 1).data(Qt.UserRole)] += \
                float(self.proxy.index(r, 6).data(Qt.UserRole))
        rows = [[k, money(v)] for k, v in sorted(sums.items())]
        rows.append(["Gesamtsumme", money(sum(sums.values()))])
        self.summary.setModel(SimpleTableModel(["Kategorie", "Zwischensumme"],
                                               rows))
        self.summary.horizontalHeader().setSectionResizeMode(
            QHeaderView.Stretch)


class SimpleTableModel(QAbstractTableModel):
    def __init__(self, headers, rows):
        super().__init__()
        self.headers, self.rows = headers, rows

    def rowCount(self, p=QModelIndex()):
        return len(self.rows)

    def columnCount(self, p=QModelIndex()):
        return len(self.headers)

    def data(self, i, role=Qt.DisplayRole):
        return self.rows[i.row()][i.column()] if i.isValid() and \
            role == Qt.DisplayRole else None

    def headerData(self, s, o, role=Qt.DisplayRole):
        return self.headers[s] if role == Qt.DisplayRole and \
            o == Qt.Horizontal else None


class SettingsTab(QWidget):
    changed = Signal()

    def __init__(self, store):
        super().__init__()
        self.store = store
        root = QHBoxLayout(self)
        self.cat = ListEditor(
            "Kategorien",
            lambda: self.store.categories,
            False,
            allow_sort=True
            )

        self.desc = ListEditor(
            "Beschreibungen",
            lambda: self.store.descriptions,
            False,
            allow_sort=True
            )

        self.acc = ListEditor(
            "Konten",
            lambda: self.store.accounts,
            True,
            allow_reorder=True
            )

        for w in (self.cat, self.desc, self.acc):
            root.addWidget(w)
            w.changed.connect(self.changed)

    def refresh(self):
        self.cat.refresh()
        self.desc.refresh()
        self.acc.refresh()


class ListEditor(QGroupBox):
    changed = Signal()

    def __init__(self, title, getter, accounts, allow_sort=False,
                 allow_reorder=False):
        super().__init__(title)
        self.getter, self.accounts = getter, accounts
        self.allow_sort = allow_sort
        self.allow_reorder = allow_reorder
        lay = QVBoxLayout(self)
        self.list = QListWidget()
        lay.addWidget(self.list)

        row = QHBoxLayout()

        add = QPushButton("Hinzufügen")
        edit = QPushButton("Bearbeiten")
        delete = QPushButton("Löschen")

        row.addWidget(add)
        row.addWidget(edit)
        row.addWidget(delete)

        if self.allow_sort:
            sort_btn = QPushButton("Sortieren A-Z")
            row.addWidget(sort_btn)
            sort_btn.clicked.connect(self.sort_items)

        if self.allow_reorder:
            up_btn = QPushButton("↑ Nach oben")
            down_btn = QPushButton("↓ Nach unten")

            row.addWidget(up_btn)
            row.addWidget(down_btn)

            up_btn.clicked.connect(self.move_up)
            down_btn.clicked.connect(self.move_down)

        lay.addLayout(row)

        add.clicked.connect(self.add)
        edit.clicked.connect(self.edit)
        delete.clicked.connect(self.delete)
        self.refresh()

    def refresh(self):
        self.list.clear()
        for x in self.getter():
            self.list.addItem(f"{x['name']} | Anfangsstand: \
                {money(float(x.get('opening_balance', 0)))}" if self.accounts
                else str(x))

    def ask(self, current=None):
        dlg = QDialog(self)
        dlg.setWindowTitle("Eintrag")
        f = QFormLayout(dlg)
        name = QLineEdit(current.get("name", "")
                         if self.accounts and current else (current or ""))
        f.addRow("Name", name)

        balance = QDoubleSpinBox()
        balance.setRange(-999999999, 999999999)
        balance.setDecimals(2)
        balance.setSuffix(" €")

        if self.accounts:
            balance.setValue(float(current.get("opening_balance", 0))
                             if current else 0)
            f.addRow("Anfangsstand", balance)
        btn = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btn.accepted.connect(dlg.accept)
        btn.rejected.connect(dlg.reject)
        f.addRow(btn)

        if dlg.exec() and name.text().strip():
            return {"name": name.text().strip(),
                    "opening_balance": balance.value()} if self.accounts \
                        else name.text().strip()

    def add(self):
        v = self.ask()
        if v is not None:
            self.getter().append(v)
            self.refresh()
            self.changed.emit()

    def edit(self):
        r = self.list.currentRow()
        if r < 0:
            return
        old = self.getter()[r]
        v = self.ask(old)
        if v is not None:
            self.getter()[r] = v
            self.refresh()
            self.changed.emit()

    def delete(self):
        r = self.list.currentRow()
        if r >= 0:
            self.getter().pop(r)
            self.refresh()
            self.changed.emit()

    def sort_items(self):
        data = self.getter()
        if self.accounts:
            return
        data.sort(key=lambda x: x.lower())
        self.refresh()
        self.changed.emit()

    def move_up(self):
        row = self.list.currentRow()
        if row <= 0:
            return
        data = self.getter()
        data[row - 1], data[row] = (
            data[row],
            data[row - 1]
            )
        self.refresh()
        self.list.setCurrentRow(row - 1)
        self.changed.emit()

    def move_down(self):
        row = self.list.currentRow()
        data = self.getter()
        if row < 0 or row >= len(data) - 1:
            return
        data[row + 1], data[row] = (
            data[row],
            data[row + 1]
            )
        self.refresh()
        self.list.setCurrentRow(row + 1)
        self.changed.emit()


class OverviewTab(QWidget):
    def __init__(self, store):
        super().__init__()
        self.store = store
        root = QVBoxLayout(self)
        top = QHBoxLayout()
        self.view = QComboBox()
        self.view.addItems(["Monatliche Ein-/Ausgänge",
                            "Regelmäßig / unregelmäßig",
                            "Gesamtkontostand",
                            "Kontenübersicht"])
        self.year = QSpinBox()
        self.year.setRange(2000, 2200)
        self.year.setValue(date.today().year)
        top.addWidget(QLabel("Ansicht"))
        top.addWidget(self.view)
        top.addWidget(QLabel("Jahr"))
        top.addWidget(self.year)
        top.addStretch()
        root.addLayout(top)
        split = QSplitter(Qt.Vertical)
        self.table = QTableView()
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        split.addWidget(self.table)
        self.chart = QChartView() if CHARTS_AVAILABLE else \
            QLabel("QtCharts ist nicht verfügbar.")
        if CHARTS_AVAILABLE:
            self.chart.setRenderHint(QPainter.Antialiasing)
        split.addWidget(self.chart)
        root.addWidget(split)
        self.view.currentIndexChanged.connect(self.refresh)
        self.year.valueChanged.connect(self.refresh)
        self.refresh()

    def refresh(self):
        idx = self.view.currentIndex()
        y = self.year.value()
        # Filter out "Umbuchung" entries
        relevant = [b for b in self.store.bookings if
                    b.beschreibung.strip().lower() != "umbuchung" and
                    int(b.datum[:4]) == y]

        if idx == 0:
            self._show_monthly_income_expenses(relevant, y)
        elif idx == 1:
            self._show_regular_irregular(relevant, y)
        elif idx == 2:
            self._show_total_balance(relevant, y)
        else:
            self._show_accounts_overview(y)

    def _show_monthly_income_expenses(self, bookings, year):
        """Monatliche Ein-/Ausgänge: Balkendiagramm, je Monat ein
            Balken für Ein- und Ausgänge"""
        income = defaultdict(float)
        expenses = defaultdict(float)

        for b in bookings:
            month = int(b.datum[5:7])
            if b.betrag >= 0:
                income[month] += b.betrag
            else:
                expenses[month] += abs(b.betrag)

        labels = [f"{m:02d}/{year}" for m in range(1, 13)]
        income_vals = [income[m] for m in range(1, 13)]
        expense_vals = [expenses[m] for m in range(1, 13)]

        # Table data
        rows = []
        for m in range(1, 13):
            rows.append([f"{m:02d}/{year}", money(income[m]),
                         money(expenses[m]),
                         money(income[m] - expenses[m])])
        rows.append(["Gesamt", money(sum(income_vals)),
                     money(sum(expense_vals)),
                     money(sum(income_vals) - sum(expense_vals))])
        self.table.setModel(SimpleTableModel(["Monat", "Eingänge",
                                              "Ausgänge", "Saldo"], rows))

        if not CHARTS_AVAILABLE:
            return

        c = QChart()
        c.setTitle("Monatliche Ein-/Ausgänge")

        # Income bar set (positive)
        income_set = QBarSet("Eingänge")
        income_set.setColor(QColor("#0b7a38"))
        income_set.append(income_vals)

        # Expenses bar set (negative, but we show as positive bars)
        expense_set = QBarSet("Ausgänge")
        expense_set.setColor(QColor("#b42318"))
        expense_set.append(expense_vals)

        series = QBarSeries()
        series.append(income_set)
        series.append(expense_set)
        c.addSeries(series)

        # X axis - months
        ax_x = QBarCategoryAxis()
        ax_x.append(labels)
        c.addAxis(ax_x, Qt.AlignBottom)
        series.attachAxis(ax_x)

        # Y axis - amounts
        ax_y = QValueAxis()
        ax_y.setTitleText("Betrag (€)")
        c.addAxis(ax_y, Qt.AlignLeft)
        series.attachAxis(ax_y)

        c.legend().setVisible(True)
        c.legend().setAlignment(Qt.AlignBottom)
        self.chart.setChart(c)

    def _show_regular_irregular(self, bookings, year):
        """Regelmäßig / unregelmäßig: Gestapelte Balkendiagramme
        mit Trennung von Ein- und Ausgaben."""
        # Separate income (>=0) and expenses (<0) for regular/irregular
        reg_income = defaultdict(float)
        reg_expense = defaultdict(float)
        irr_income = defaultdict(float)
        irr_expense = defaultdict(float)

        for b in bookings:
            month = int(b.datum[5:7])
            if b.regelmaessig:
                if b.betrag >= 0:
                    reg_income[month] += b.betrag
                else:
                    reg_expense[month] += b.betrag
            else:
                if b.betrag >= 0:
                    irr_income[month] += b.betrag
                else:
                    irr_expense[month] += b.betrag

        labels = [f"{m:02d}/{year}" for m in range(1, 13)]
        reg_income_vals = [reg_income[m] for m in range(1, 13)]
        reg_expense_vals = [reg_expense[m] for m in range(1, 13)]
        irr_income_vals = [irr_income[m] for m in range(1, 13)]
        irr_expense_vals = [irr_expense[m] for m in range(1, 13)]

        # Table data
        rows = []
        for m in range(1, 13):
            reg_total = reg_income[m] + reg_expense[m]
            irr_total = irr_income[m] + irr_expense[m]
            rows.append([f"{m:02d}/{year}",
                         money(reg_income[m]), money(reg_expense[m]),
                         money(reg_total),
                         money(irr_income[m]), money(irr_expense[m]),
                         money(irr_total)])
        rows.append(["Gesamt",
                     money(sum(reg_income_vals)), money(sum(reg_expense_vals)),
                     money(sum(reg_income_vals) + sum(reg_expense_vals)),
                     money(sum(irr_income_vals)), money(sum(irr_expense_vals)),
                     money(sum(irr_income_vals) + sum(irr_expense_vals))])
        self.table.setModel(SimpleTableModel(
            ["Monat", "Reg. Einnahmen", "Reg. Ausgaben", "Reg. Summe",
             "Unreg. Einnahmen", "Unreg. Ausgaben", "Unreg. Summe"], rows))

        if not CHARTS_AVAILABLE:
            return

        c = QChart()
        c.setTitle("Regelmäßig / Unregelmäßig (Ein- & Ausgaben)")

        # Regular income (positive, blue)
        reg_inc_set = QBarSet("Reg. Einnahmen")
        reg_inc_set.setColor(QColor("#1565C0"))
        reg_inc_set.append(reg_income_vals)

        # Regular expenses (negative, light blue)
        reg_exp_set = QBarSet("Reg. Ausgaben")
        reg_exp_set.setColor(QColor("#64B5F6"))
        reg_exp_set.append(reg_expense_vals)

        # Irregular income (positive, orange)
        irr_inc_set = QBarSet("Unreg. Einnahmen")
        irr_inc_set.setColor(QColor("#E65100"))
        irr_inc_set.append(irr_income_vals)

        # Irregular expenses (negative, light orange)
        irr_exp_set = QBarSet("Unreg. Ausgaben")
        irr_exp_set.setColor(QColor("#FFB74D"))
        irr_exp_set.append(irr_expense_vals)

        # Stacked bar series: income stacks above 0, expenses below 0
        series = QStackedBarSeries()
        series.append(reg_inc_set)
        series.append(reg_exp_set)
        series.append(irr_inc_set)
        series.append(irr_exp_set)
        c.addSeries(series)

        # X axis - months
        ax_x = QBarCategoryAxis()
        ax_x.append(labels)
        c.addAxis(ax_x, Qt.AlignBottom)
        series.attachAxis(ax_x)

        # Y axis - amounts
        ax_y = QValueAxis()
        ax_y.setTitleText("Betrag (€)")
        c.addAxis(ax_y, Qt.AlignLeft)
        series.attachAxis(ax_y)

        c.legend().setVisible(True)
        c.legend().setAlignment(Qt.AlignBottom)
        self.chart.setChart(c)

    def _show_total_balance(self, bookings, year):
        """Gesamtkontostand: Liniendiagramm mit monatlichen Werten"""
        # Calculate opening balance
        base = sum(float(a.get("opening_balance", 0)) for a
                   in self.store.accounts)

        # Aggregate bookings by month
        monthly_changes = defaultdict(float)
        for b in bookings:
            month = int(b.datum[5:7])
            monthly_changes[month] += b.betrag

        # Calculate running balance per month
        labels = [f"{m:02d}/{year}" for m in range(1, 13)]
        running = base
        vals = []
        for m in range(1, 13):
            running += monthly_changes[m]
            vals.append(running)

        # Table data
        rows = [[f"{m:02d}/{year}", money(vals[m-1])] for m in range(1, 13)]
        self.table.setModel(SimpleTableModel(["Monat", "Kontostand"], rows))

        if not CHARTS_AVAILABLE:
            return

        c = QChart()
        c.setTitle("Gesamtkontostand")

        s = QLineSeries()
        pen = QPen(QColor("#2196F3"))
        pen.setWidth(2)
        s.setPen(pen)
        for i, v in enumerate(vals):
            s.append(i, v)
        c.addSeries(s)

        # X axis - months as categories
        ax_x = QBarCategoryAxis()
        ax_x.append(labels)
        c.addAxis(ax_x, Qt.AlignBottom)
        s.attachAxis(ax_x)

        # Y axis - balance
        ax_y = QValueAxis()
        ax_y.setTitleText("Kontostand (€)")
        c.addAxis(ax_y, Qt.AlignLeft)
        s.attachAxis(ax_y)

        c.legend().setVisible(False)
        self.chart.setChart(c)

    def _show_accounts_overview(self, year):
        rows = []
        total = 0
        for a in self.store.accounts:
            start = float(a.get("opening_balance", 0))
            end = start + sum(b.betrag for b in self.store.bookings if
                              b.konto == a["name"] and
                              int(b.datum[:4]) <= year)
            diff = end-start
            total += end
            rows.append([a["name"], money(start), money(end), money(diff)])

        rows.append(["Summe aller Konten", "", money(total), ""])
        self.table.setModel(SimpleTableModel(["Konto",
                                              "Anfangsstand",
                                              "Endstand",
                                              "Differenz"], rows))
        self.clear_chart("Kontenübersicht")

    def clear_chart(self, title):
        if CHARTS_AVAILABLE:
            c = QChart()
            c.setTitle(title)
            self.chart.setChart(c)


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.store = Store()
        self.model = BookingModel(self.store)
        self.setWindowTitle(APP_TITLE)
        self.resize(1250, 800)
        self.tabs = QTabWidget()
        self.entry = EntryTab(self.store, self.model)
        self.pivot = PivotTab(self.store, self.model)
        self.settings = SettingsTab(self.store)
        self.overview = OverviewTab(self.store)

        for w, n in [(self.entry, "Dateneingabe"),
                     (self.pivot, "Pivot-Ansicht"),
                     (self.overview, "Übersicht / Diagramme"),
                     (self.settings, "Einstellungen")]:
            self.tabs.addTab(w, n)
        self.setCentralWidget(self.tabs)
        self.create_menu()
        self.entry.changed.connect(self.data_changed)
        self.settings.changed.connect(self.settings_changed)
        self.tabs.currentChanged.connect(self.on_tab_changed)
        self._pivot_needs_refresh = False
        self._overview_needs_refresh = False
        self._entry_needs_refresh = False
        self.statusBar().showMessage("Bereit")

    def create_menu(self):
        m = self.menuBar().addMenu("Datei")
        for text, slot, shortcut in [("Neu", self.new, "Ctrl+N"),
                                     ("JSON öffnen…", self.open_json,
                                      "Ctrl+O"),
                                     ("Speichern", self.save, "Ctrl+S"),
                                     ("Speichern unter…", self.save_as,
                                      "Ctrl+Shift+S"),
                                     ("CSV/TXT importieren…",
                                      self.import_delimited, "Ctrl+I"),
                                     ("CSV/TXT exportieren…",
                                      self.export_delimited, "Ctrl+E")]:
            a = QAction(text, self)
            a.setShortcut(shortcut)
            a.triggered.connect(slot)
            m.addAction(a)
        m.addSeparator()
        # SQLite
        sqlite_open = QAction("SQLite öffnen…", self)
        sqlite_open.triggered.connect(self.open_sqlite)
        m.addAction(sqlite_open)
        sqlite_save = QAction("Als SQLite speichern…", self)
        sqlite_save.triggered.connect(self.save_sqlite)
        m.addAction(sqlite_save)
        m.addSeparator()
        # Excel
        excel_export = QAction("Excel exportieren…", self)
        excel_export.triggered.connect(self.export_excel)
        m.addAction(excel_export)
        m.addSeparator()
        quit_a = QAction("Beenden", self)
        quit_a.triggered.connect(self.close)
        m.addAction(quit_a)

        help_menu = self.menuBar().addMenu("Hilfe")
        info_a = QAction("Info", self)
        info_a.triggered.connect(self.show_info)
        help_menu.addAction(info_a)
        modules_a = QAction("Module", self)
        modules_a.triggered.connect(self.show_modules)
        help_menu.addAction(modules_a)

    def data_changed(self):
        idx = self.tabs.currentIndex()
        if idx == 1:
            self.pivot.refresh_settings()
            self.pivot.apply()
        else:
            self._pivot_needs_refresh = True
        if idx == 2:
            self.overview.refresh()
        else:
            self._overview_needs_refresh = True
        if idx == 0:
            self.entry._refresh_balances()
        else:
            self._entry_needs_refresh = True

    def settings_changed(self):
        self.entry.refresh_settings()
        self.pivot.refresh_settings()
        idx = self.tabs.currentIndex()
        if idx == 1:
            self.pivot.apply()
        else:
            self._pivot_needs_refresh = True
        if idx == 2:
            self.overview.refresh()
        else:
            self._overview_needs_refresh = True

    def on_tab_changed(self, index):
        if index == 0:
            if self._entry_needs_refresh:
                self.entry._refresh_balances()
                self._entry_needs_refresh = False
        elif index == 1:
            if self._pivot_needs_refresh:
                self.pivot.refresh_settings()
                self.pivot.apply()
                self._pivot_needs_refresh = False
        elif index == 2:
            if self._overview_needs_refresh:
                self.overview.refresh()
                self._overview_needs_refresh = False

    def new(self):
        self.store = Store()
        self.rebind()

    def rebind(self):
        self.model.store = self.store
        self.model.refresh()
        self.entry.store = self.store
        self.pivot.store = self.store
        self.settings.store = self.store
        self.overview.store = self.store
        self.entry.refresh_settings()
        self.settings.cat.getter = lambda: self.store.categories
        self.settings.desc.getter = lambda: self.store.descriptions
        self.settings.acc.getter = lambda: self.store.accounts
        self.settings.refresh()
        self.data_changed()

    def open_json(self):
        p, _ = QFileDialog.getOpenFileName(self, "JSON laden", "",
                                           "JSON (*.json)")
        if p:
            try:
                self.store.load(Path(p))
                self.rebind()
                self.statusBar().showMessage(f"Geladen: {p}")
            except Exception as e:
                QMessageBox.critical(self, "Fehler", str(e))

    def save(self):
        if not self.store.path:
            return self.save_as()
        try:
            self.store.save()
            self.statusBar().showMessage(f"Gespeichert: {self.store.path}")
        except Exception as e:
            QMessageBox.critical(self, "Fehler", str(e))

    def save_as(self):
        p, _ = QFileDialog.getSaveFileName(self,
                                           "JSON speichern",
                                           str(self.store.path or
                                               DEFAULT_FILE),
                                           "JSON (*.json)")
        if p:
            if not p.lower().endswith('.json'):
                p += '.json'
            try:
                self.store.save(Path(p))
                self.statusBar().showMessage(f"Gespeichert: {p}")
            except Exception as e:
                QMessageBox.critical(self, "Fehler", str(e))

    def import_delimited(self):
        p, _ = QFileDialog.getOpenFileName(self,
                                           "Daten importieren",
                                           "",
                                           "CSV/TXT (*.csv *.txt)")
        if not p:
            return
        try:
            sample = Path(p).read_text(encoding="utf-8-sig")[:4096]
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
            with open(p, encoding="utf-8-sig", newline="") as f:
                for r in csv.DictReader(f, dialect=dialect):
                    self.store.bookings.append(Booking(r.get("id") or
                                                       str(uuid.uuid4()),
                                                       r["Datum"],
                                                       r["Kategorie"],
                                                       r.get("Beschreibung",
                                                             ""),
                                                       parse_bool(
                                                           r.get("PayPal",
                                                                 False)),
                                                       parse_bool(r.get(
                                                           "Regelmäßig",
                                                           False)),
                                                       r["Konto"],
                                                       float(r["Betrag"]
                                                             .replace(",",
                                                                      "."))))
            self.model.refresh()
            self.data_changed()
        except Exception as e:
            QMessageBox.critical(self,
                                 "Importfehler",
                                 f"Erwartete Spalten: Datum, Kategorie,\
                                     Beschreibung, PayPal, Regelmäßig, Konto,\
                                         Betrag\n\n{e}")

    def export_delimited(self):
        p, _ = QFileDialog.getSaveFileName(self,
                                           "Daten exportieren",
                                           "buchungen.csv",
                                           "CSV (*.csv);;TXT (*.txt)")
        if not p:
            return
        delim = '\t' if p.lower().endswith('.txt') else ';'
        try:
            with open(p, 'w', encoding='utf-8-sig', newline='') as f:
                w = csv.writer(f, delimiter=delim)
                w.writerow(["id", "Datum", "Kategorie", "Beschreibung",
                            "PayPal", "Regelmäßig", "Konto", "Betrag"])
                for b in self.store.bookings:
                    w.writerow([b.id,
                                b.datum,
                                b.kategorie,
                                b.beschreibung,
                                "Ja" if b.paypal else "Nein",
                                "Ja" if b.regelmaessig else "Nein",
                                b.konto,
                                str(b.betrag).replace('.', ',')])
            self.statusBar().showMessage(f"Exportiert: {p}")
        except Exception as e:
            QMessageBox.critical(self, "Exportfehler", str(e))

    def open_sqlite(self):
        p, _ = QFileDialog.getOpenFileName(self, "SQLite laden", "",
                                           "SQLite (*.db *.sqlite *.sqlite3)")
        if p:
            try:
                self.store.load_sqlite(Path(p))
                self.rebind()
                self.statusBar().showMessage(f"Geladen: {p}")
            except Exception as e:
                QMessageBox.critical(self, "Fehler", str(e))

    def save_sqlite(self):
        p, _ = QFileDialog.getSaveFileName(self, "SQLite speichern",
                                           str(self.store.path or
                                               "haushaltsbuch.db"),
                                           "SQLite (*.db *.sqlite *.sqlite3)")
        if p:
            if not p.lower().endswith(('.db', '.sqlite', '.sqlite3')):
                p += '.db'
            try:
                self.store.save_sqlite(Path(p))
                self.statusBar().showMessage(f"Gespeichert: {p}")
            except Exception as e:
                QMessageBox.critical(self, "Fehler", str(e))

    def export_excel(self):
        if not EXCEL_AVAILABLE:
            QMessageBox.critical(self, "Fehler",
                                 "openpyxl ist nicht installiert.\n"
                                 "pip install openpyxl")
            return
        p, _ = QFileDialog.getSaveFileName(self, "Excel exportieren",
                                           "buchungen.xlsx",
                                           "Excel (*.xlsx)")
        if p:
            if not p.lower().endswith('.xlsx'):
                p += '.xlsx'
            try:
                self.store.export_excel(Path(p))
                self.statusBar().showMessage(f"Exportiert: {p}")
            except Exception as e:
                QMessageBox.critical(self, "Exportfehler", str(e))

    def show_info(self):
        QMessageBox.information(
            self, "Info",
            f"""Haushaltsbuch
Version: {VERSION}
Copyright (C) 2026 {DEVELOPER}

This program is free software: you can redistribute it and/or modify
it under the terms of the GNU General Public License version 3 of the
License.

This program is distributed in the hope that it will be useful,
but WITHOUT ANY WARRANTY; without even the implied warranty of
MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
GNU General Public License for more details.

You should have received a copy of the GNU General Public License
along with this program.  If not, see <https://www.gnu.org/licenses/>.""")

    def show_modules(self):
        lines = ["Genutzte Module:", ""]
        lines.append(f"Python: {sys.version.split()[0]}")
        try:
            import PySide6
            lines.append(f"PySide6: {PySide6.__version__}")
        except Exception:
            lines.append("PySide6: nicht verfügbar")
        lines.append(f"QtCharts: {'verfügbar' if CHARTS_AVAILABLE else
                     'nicht verfügbar'}")
        if EXCEL_AVAILABLE:
            try:
                import openpyxl
                lines.append(f"openpyxl: {openpyxl.__version__}")
            except Exception:
                lines.append("openpyxl: verfügbar")
        else:
            lines.append("openpyxl: nicht installiert")
        lines.append(f"sqlite3: {sqlite3.sqlite_version}")
        lines.append("csv, json, shutil, uuid, datetime, collections, "
                     "dataclasses, pathlib: Standardbibliothek")

        dlg = QDialog(self)
        dlg.setWindowTitle("Module")
        dlg.resize(420, 320)
        layout = QVBoxLayout(dlg)
        text = QTextEdit()
        text.setReadOnly(True)
        text.setPlainText("\n".join(lines))
        layout.addWidget(text)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok)
        buttons.accepted.connect(dlg.accept)
        layout.addWidget(buttons)
        dlg.exec()

    # Aliases
    def info(self):
        return self.show_info()

    def modules(self):
        return self.show_modules()


def main():
    app = QApplication(sys.argv)
    app.setApplicationName(APP_TITLE)
    app.setStyle(QStyleFactory.create("Fusion"))
    app.setWindowIcon(QIcon(
        "C:/Users/awind/OneDrive/Dokumente/GitHub/Haushaltsbuch/icons/icon.svg"))
    QGuiApplication.styleHints().setColorScheme(Qt.ColorScheme.Light)
    w = MainWindow()
    w.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
